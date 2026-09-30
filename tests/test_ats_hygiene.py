import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / ".claude" / "skills" / "resume-tailor"
sys.path.insert(0, str(SKILL / "scripts"))

import ats_hygiene as H  # noqa: E402
from packet_lint import check_build_script  # noqa: E402

CAND = {"name": "Jordan Q Example", "location": "City, ST", "phone": "(555) 010-0000",
        "email": "jordan@example.com", "linkedin": "linkedin.com/in/jordan-example"}


def _doc_with_bullet():
    doc = Document()
    doc.add_paragraph("Jordan Q Example")
    doc.add_paragraph("Triaged alerts across 100+ client environments.", style="List Bullet")
    return doc


def test_raw_python_docx_save_is_flagged(tmp_path):
    path = tmp_path / "resume.docx"
    _doc_with_bullet().save(path)
    problems = " | ".join(H.verify_docx(path))
    assert "python-docx" in problems
    assert "U+F0B7" in problems
    assert "thumbnail" in problems
    assert "created date" in problems


def test_write_document_leaves_no_generator_fingerprint(tmp_path):
    out = H.write_document(_doc_with_bullet(), tmp_path / "resume.docx", CAND, "resume", pdf=False)
    assert H.verify_docx(out["docx"]) == []

    doc = Document(out["docx"])
    cp = doc.core_properties
    assert cp.author == cp.last_modified_by == "Jordan Q Example"
    assert cp.title == "Jordan Q Example Resume"
    assert cp.comments == ""
    assert cp.created.year == datetime.now(timezone.utc).year

    numbering = doc.part.numbering_part.element.xml
    assert H.PUA_BULLET not in numbering and H.BULLET in numbering
    app = next(p for p in doc.part.package.iter_parts() if str(p.partname) == "/docProps/app.xml")
    assert b"<Words>0</Words>" not in app.blob


def test_cover_letter_kind_and_bad_kind(tmp_path):
    out = H.write_document(Document(), tmp_path / "cover_letter.docx", CAND, "cover_letter", pdf=False)
    assert Document(out["docx"]).core_properties.title == "Jordan Q Example Cover Letter"
    with pytest.raises(ValueError):
        H.write_document(Document(), tmp_path / "memo.docx", CAND, "memo", pdf=False)


def test_upload_names():
    assert H.upload_name(CAND, "resume") == "Jordan_Example_Resume.pdf"
    assert H.upload_name(CAND, "cover_letter") == "Jordan_Example_Cover_Letter.pdf"
    assert H.upload_name({**CAND, "file_stem": "JExample"}, "resume") == "JExample_Resume.pdf"
    assert H.upload_name({"name": "Cher"}, "resume") == "Cher_Resume.pdf"


def test_verify_pdf_finds_marker(tmp_path):
    dirty = tmp_path / "dirty.pdf"
    dirty.write_bytes(b"%PDF-1.6\n1 0 obj << /Author (python-docx) >> endobj\n%%EOF\n")
    clean = tmp_path / "clean.pdf"
    clean.write_bytes(b"%PDF-1.6\n1 0 obj << /Author (Jordan Q Example) >> endobj\n%%EOF\n")
    assert H.verify_pdf(dirty)
    assert H.verify_pdf(clean) == []
    assert H.verify_pdf(tmp_path / "missing.pdf") == []


def test_template_writes_clean_documents(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "master_bullets.yaml").write_text(
        "candidate:\n"
        "  name: Jordan Q Example\n  location: City, ST\n  phone: (555) 010-0000\n"
        "  email: jordan@example.com\n  linkedin: linkedin.com/in/jordan-example\n",
        encoding="utf-8")
    shutil.copytree(SKILL / "scripts", tmp_path / ".claude" / "skills" / "resume-tailor" / "scripts")
    pkt = tmp_path / "output" / "Acme_1"
    pkt.mkdir(parents=True)
    shutil.copy(SKILL / "templates" / "build_docs.py", pkt / "build_docs.py")

    assert check_build_script(pkt / "build_docs.py") == []
    subprocess.run([sys.executable, str(pkt / "build_docs.py"), "--no-pdf"], cwd=tmp_path, check=True,
                   capture_output=True)
    for name in ("resume.docx", "cover_letter.docx"):
        assert H.verify_docx(pkt / name) == []
