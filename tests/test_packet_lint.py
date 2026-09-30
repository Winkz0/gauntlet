import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / ".claude" / "skills" / "resume-tailor" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import packet_lint as L  # noqa: E402
from ats_hygiene import write_document  # noqa: E402

LIBRARY = """\
candidate:
  name: Jordan Q Example
  location: City, ST
  phone: (555) 010-0000
  email: jordan@example.com
  linkedin: linkedin.com/in/jordan-example
bullets:
  - text: "Lead investigation and case management for incidents escalated from Tier I, with a 99% quality score."
    tools: [CrowdStrike Falcon]
    metric: "99% quality score"
  - text: "Primary escalation point for a 15+ analyst Tier I team."
    tools: []
projects:
  - name: Lab
    text: "Built a malware lab."
    tools: [Python, REMnux]
tool_inventory:
  edr_siem: [CrowdStrike Falcon, Splunk]
  familiarity: [Active Directory]
"""

POSTING = (
    "What you'll do: own the full detection efficacy review lifecycle for customer escalations. "
    "Lead investigation and case management for incidents escalated from Tier I. "
    "Requirements: CrowdStrike, Python, Active Directory, and Azure. "
    "Bachelor's degree in Computer Science or equivalent experience. "
    "Education: BA or BS degree in Information Security. "
)


@pytest.fixture
def repo(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "master_bullets.yaml").write_text(LIBRARY, encoding="utf-8")
    (tmp_path / "posting.txt").write_text(POSTING, encoding="utf-8")
    return tmp_path


def make_packet(repo, bullets, skills=(), clean=True, name="Acme_7"):
    pkt = repo / "output" / name
    pkt.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("Jordan Q Example")
    doc.add_paragraph("City, ST | (555) 010-0000 | jordan@example.com")
    doc.add_paragraph("SOC Analyst II | Sep 2025 - Present")
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")
    for label, items in skills:
        doc.add_paragraph(f"{label}: {items}")
    cand = {"name": "Jordan Q Example"}
    if clean:
        write_document(doc, pkt / "resume.docx", cand, "resume", pdf=False)
    else:
        doc.save(pkt / "resume.docx")
    return pkt


def test_numbers_must_trace_to_library(repo):
    pkt = make_packet(repo, [
        "Lead investigation and case management with a 99% quality score.",
        "Cut alert volume by 40% across 15+ analysts.",
        "Matched SHA256 hashes to samples.",
    ])
    report = L.lint_packet(repo, pkt, POSTING)
    assert report["fact_trace"]["numbers_not_in_library"] == ["resume: 40%"]


def test_skills_lines_must_trace(repo):
    pkt = make_packet(repo, ["Built a malware lab."],
                      skills=[("Endpoint", "CrowdStrike Falcon; Ghidra"), ("Identity", "Active Directory")])
    ft = L.lint_packet(repo, pkt, POSTING)["fact_trace"]
    assert ft["tools_not_in_library"] == ["Ghidra"]
    assert ft["familiarity_tier_used"] == ["Active Directory"]


def test_coverage_buckets(repo):
    pkt = make_packet(repo, ["Worked CrowdStrike Falcon detections."])
    cov = L.lint_packet(repo, pkt, POSTING)["coverage"]
    assert "crowdstrike" in cov["on_resume"]
    assert "python" in cov["library_not_on_resume"]
    assert cov["familiarity_tier_only"] == ["active directory"]
    assert "azure" in cov["not_in_library"]


def test_phrase_lift_ignores_the_candidates_own_wording(repo):
    pkt = make_packet(repo, [
        "Lead investigation and case management for incidents escalated from Tier I.",
        "Own the full detection efficacy review lifecycle for customers.",
    ])
    lifted = L.lint_packet(repo, pkt, POSTING)["phrase_lift"]["resume"]
    assert any("detection efficacy review lifecycle" in g for g in lifted)
    assert not any("case management for incidents" in g for g in lifted)


def test_knockouts_read_equivalency():
    hits = L.check_knockouts(POSTING)
    degree = [h for h in hits if "degree" in h["text"].lower() or "bachelor" in h["match"].lower()]
    assert any(h["equivalency_clause"] for h in degree)
    assert any(not h["equivalency_clause"] and "BA or BS" in h["text"] for h in degree)


def test_build_script_check_reads_code_not_comments(tmp_path):
    good = tmp_path / "good.py"
    good.write_text('"""Never call doc.save() here."""\nwrite_document(doc, p, c, "resume")\n')
    bad = tmp_path / "bad.py"
    bad.write_text("doc.save(path)\n")
    assert L.check_build_script(good) == []
    assert len(L.check_build_script(bad)) == 2


def test_staleness(repo):
    pkt = make_packet(repo, ["Built a malware lab."])
    sha = L.load_library(repo)["sha"]
    assert L.staleness(pkt, sha).startswith("unrecorded")
    (pkt / "gauntlet.json").write_text(json.dumps({"library_sha": sha}))
    assert L.staleness(pkt, sha) == "current"
    (pkt / "gauntlet.json").write_text(json.dumps({"library_sha": "000000000000"}))
    assert L.staleness(pkt, sha).startswith("STALE")


def _cli(repo, *args):
    return subprocess.run([sys.executable, str(SCRIPTS / "packet_lint.py"), *args, "--root", str(repo)],
                          capture_output=True, text=True)


def test_strict_gate(repo):
    make_packet(repo, ["Built a malware lab."], clean=False, name="Dirty_1")
    make_packet(repo, ["Built a malware lab."], clean=True, name="Clean_2")
    dirty = _cli(repo, "output/Dirty_1", "--posting", str(repo / "posting.txt"), "--strict")
    clean = _cli(repo, "output/Clean_2", "--posting", str(repo / "posting.txt"), "--strict")
    assert dirty.returncode == 1 and "BLOCKED" in dirty.stdout
    assert clean.returncode == 0 and "GATE: pass" in clean.stdout


def test_predraft_and_brief_modes(repo):
    pre = _cli(repo, "--posting", str(repo / "posting.txt"), "--json")
    out = json.loads(pre.stdout)
    assert "on_resume" not in out["coverage"] and "python" in out["coverage"]["library_not_on_resume"]
    make_packet(repo, ["Built a malware lab."], name="Brief_3")
    brief = _cli(repo, "output/Brief_3", "--posting", str(repo / "posting.txt"), "--review-brief")
    assert "=== POSTING" in brief.stdout and "Built a malware lab." in brief.stdout
    assert "master_bullets" not in brief.stdout
