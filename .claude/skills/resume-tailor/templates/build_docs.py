#!/usr/bin/env python3
"""
Build the tailored packet for job <JOB_ID>:
<Company>, <Title> (<requisition, if any>).

    python output/<Company>_<JOB_ID>/build_docs.py

Template: copy to output/<Company>_<JOB_ID>/build_docs.py and edit only the
CONTENT section. Name and contact come from data/master_bullets.yaml; every
claim in CONTENT must trace to a bullet, project, certification, or tool
entry in that file.

Every document is written through ats_hygiene.write_document(), which cleans
the generator metadata and bullet glyphs, saves the .docx, verifies it,
renders the PDF with LibreOffice (soffice --headless), verifies that, and
writes an upload-named copy. Never call doc.save() here.
"""
from __future__ import annotations

import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "resume-tailor" / "scripts"))
from ats_hygiene import write_document  # noqa: E402

BULLETS = yaml.safe_load((ROOT / "data" / "master_bullets.yaml").read_text(encoding="utf-8"))
CAND = BULLETS["candidate"]
CONTACT = " | ".join([CAND["location"], CAND["phone"], CAND["email"], CAND["linkedin"]])

# ================================================================ CONTENT
COMPANY = "<Company>"

SUMMARY = "<Two or three sentences, selected from professional_summary / summary_flex_angles.>"

EXPERIENCE = [
    {
        "company": "<Employer>", "location": "<Location>",
        "roles": [{
            "title": "<Title>", "dates": "<Mon YYYY - Mon YYYY>",
            "bullets": [
                "<Bullet selected from master_bullets.yaml.>",
            ],
        }],
    },
]

PROJECTS = [
    ("<Project name> (personal project)", "<Project text from master_bullets.yaml projects.>"),
]

SKILLS_HEADING = "Certifications and Skills"
SKILLS = [
    ("Certifications", "<From master_bullets.yaml certifications.>"),
    ("<Label>", "<Tools from tool_inventory; never the familiarity tier without confirmation.>"),
]

LETTER_DATE = "<Month D, YYYY>"
LETTER_RE = "Re: <Title> (<requisition>)"
LETTER = [
    f"Dear {COMPANY} Hiring Team,",
    "<Three to four short paragraphs.>",
    "Thank you,",
    CAND["name"],
]
# ============================================================ END CONTENT


# ---------------------------------------------------------------- helpers
def _base_doc() -> Document:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.5)
        s.left_margin = s.right_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)
    return doc


def _para(doc, text="", bold=False, size=None, align=None, after=None, before=None):
    p = doc.add_paragraph()
    if text:
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    return p


def _heading(doc, text):
    _para(doc, text.upper(), bold=True, size=11, before=6, after=1)


def _bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(0)
    return p


def _header(doc):
    _para(doc, CAND["name"], bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    _para(doc, CONTACT, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)


# ---------------------------------------------------------------- builders
def build_resume() -> Document:
    doc = _base_doc()
    _header(doc)

    _heading(doc, "Summary")
    _para(doc, SUMMARY)

    _heading(doc, "Experience")
    for job in EXPERIENCE:
        _para(doc, f"{job['company']} | {job['location']}", bold=True, before=4, after=0)
        for role in job["roles"]:
            p = _para(doc, after=1)
            p.add_run(role["title"]).italic = True
            p.add_run(f" | {role['dates']}")
            for b in role["bullets"]:
                _bullet(doc, b)

    if PROJECTS:
        _heading(doc, "Projects")
        for title, text in PROJECTS:
            _para(doc, title, bold=True, before=2, after=0)
            _bullet(doc, text)

    _heading(doc, SKILLS_HEADING)
    for label, items in SKILLS:
        p = _para(doc, after=1)
        p.add_run(f"{label}: ").bold = True
        p.add_run(items)
    return doc


def build_cover() -> Document:
    doc = _base_doc()
    doc.styles["Normal"].font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    _header(doc)
    _para(doc, LETTER_DATE, before=6, after=6)
    _para(doc, f"{COMPANY} Hiring Team", after=0)
    _para(doc, LETTER_RE, after=10)
    for i, block in enumerate(LETTER):
        last = i >= len(LETTER) - 2
        _para(doc, block, after=0 if last else 8)
    return doc


def main(pdf: bool = True) -> None:
    write_document(build_resume(), HERE / "resume.docx", CAND, "resume", pdf=pdf)
    write_document(build_cover(), HERE / "cover_letter.docx", CAND, "cover_letter", pdf=pdf)


if __name__ == "__main__":
    main(pdf="--no-pdf" not in sys.argv)
